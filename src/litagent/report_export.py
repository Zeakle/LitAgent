"""Render immutable report versions to local Markdown, PDF, and Word files."""

from __future__ import annotations

import asyncio
import html
import io
import re
import sys
from typing import Literal

import fitz
from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from markdown_it import MarkdownIt

from litagent.citation_formats import STYLES
from litagent.report_render import (
    ReaderReport,
    _without_managed_markers,
    render_report_html,
)
from litagent.workbench import ExportResult, ReportVersion

_DRAFT_LABEL = "DRAFT - NOT APPROVED FOR PUBLICATION"


class LocalReportExporter:
    """Render server-owned report text without model or network access."""

    async def export(
        self,
        report: ReportVersion,
        *,
        format: Literal["md", "pdf", "docx"],
        publishable: bool,
        citation_style: str = "evidence",
        reader_report: ReaderReport | None = None,
    ) -> ExportResult:
        """Render one exact revision and expose its source-content hash."""
        if format not in {"md", "pdf", "docx"}:
            raise ValueError("unsupported_export_format")
        if citation_style not in STYLES:
            raise ValueError("unsupported_citation_style")
        if citation_style != "evidence" and (
            reader_report is None
            or reader_report.source_sha256 != report.content_sha256
        ):
            raise ValueError("reader_projection_unavailable")
        markdown = (
            reader_report.markdown
            if citation_style != "evidence" and reader_report is not None
            else report.survey
        )
        if (
            format != "md"
            and reader_report is not None
            and citation_style != "evidence"
        ):
            markdown = _export_reader_markdown(reader_report)
        content = await asyncio.to_thread(
            self._render, report, markdown, format, publishable
        )
        media_type = {
            "md": "text/markdown; charset=utf-8",
            "pdf": "application/pdf",
            "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        }[format]
        return ExportResult(
            filename=f"litagent-{report.run_id}-r{report.revision}.{format}",
            media_type=media_type,
            content=content,
            report_sha256=report.content_sha256,
        )

    def _render(
        self, report: ReportVersion, markdown: str, format: str, publishable: bool
    ) -> bytes:
        if format == "md":
            return self._markdown(report, markdown, publishable).encode("utf-8")
        markdown = _without_managed_markers(markdown)
        if format == "pdf":
            return self._pdf(report, markdown, publishable)
        return self._docx(report, markdown, publishable)

    @staticmethod
    def _metadata_lines(report: ReportVersion, publishable: bool) -> list[str]:
        lines = [] if publishable else [_DRAFT_LABEL]
        lines.extend(
            [
                f"Run: {report.run_id}",
                f"Revision {report.revision}",
                f"Language: {report.report_language}",
                f"Human review: {report.human_status}",
                f"Content SHA-256: {report.content_sha256}",
            ]
        )
        return lines

    def _markdown(self, report: ReportVersion, markdown: str, publishable: bool) -> str:
        metadata = "\n".join(
            f"> {line}" for line in self._metadata_lines(report, publishable)
        )
        return f"{metadata}\n\n{markdown}\n"

    def _pdf(self, report: ReportVersion, markdown: str, publishable: bool) -> bytes:
        # Story receives escaped HTML without images, so it never fetches
        # untrusted assets embedded by a report or follows links at render time.
        rendered = render_report_html(markdown).html
        rendered = re.sub(r"<img\b[^>]*>", "[image omitted]", rendered)
        metadata = "<br>".join(
            html.escape(line) for line in self._metadata_lines(report, publishable)
        )
        css = (
            "body{font-family:sans-serif;font-size:10pt;line-height:1.4}"
            "h1{font-size:19pt}h2{font-size:15pt}h3{font-size:12pt}"
            "table{border-collapse:collapse;width:100%}"
            "td,th{border:1px solid #888;padding:4pt}"
            "pre{white-space:pre-wrap}a{color:#075c83}"
        )
        story = fitz.Story(
            html=f"<p class='metadata'>{metadata}</p>{rendered}", user_css=css
        )
        page_rect = fitz.paper_rect("a4")
        document = story.write_with_links(
            lambda _page, _filled: (
                page_rect,
                fitz.Rect(48, 48, page_rect.width - 48, page_rect.height - 48),
                None,
            ),
        )
        try:
            if not publishable:
                # Keep the warning visible after page one without dropping links.
                for page in document:
                    page.insert_text(
                        (48, page.rect.height - 30),
                        _DRAFT_LABEL,
                        fontsize=8,
                        color=(0.56, 0.25, 0.15),
                    )
            return document.tobytes()
        finally:
            document.close()

    def _docx(self, report: ReportVersion, markdown: str, publishable: bool) -> bytes:
        document = Document()
        # Explicit East Asian fonts prevent theme fallbacks from dropping CJK glyphs.
        east_asian_font = (
            "Arial Unicode MS" if sys.platform == "darwin" else "Microsoft YaHei"
        )
        for style in document.styles:
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:eastAsia"), east_asian_font)
            fonts.attrib.pop(qn("w:eastAsiaTheme"), None)
        document.sections[0].header.paragraphs[0].text = (
            _DRAFT_LABEL if not publishable else ""
        )
        for line in self._metadata_lines(report, publishable):
            document.add_paragraph(line, style="Subtitle")
        tokens = (
            MarkdownIt("commonmark", {"html": False}).enable("table").parse(markdown)
        )
        list_stack: list[str] = []
        paragraph = None
        table = None
        row = None
        cell = None
        column = 0
        for index, token in enumerate(tokens):
            kind = token.type
            if kind in {"bullet_list_open", "ordered_list_open"}:
                list_stack.append(kind)
            elif kind in {"bullet_list_close", "ordered_list_close"}:
                list_stack.pop()
            elif kind == "table_open":
                columns = 0
                for entry in tokens[index + 1 :]:
                    if entry.type == "tr_close":
                        break
                    columns += entry.type in {"th_open", "td_open"}
                table = document.add_table(rows=0, cols=max(columns, 1))
                table.style = "Table Grid"
            elif kind == "table_close":
                table = None
            elif kind == "tr_open" and table is not None:
                row = table.add_row()
                column = 0
            elif kind in {"th_open", "td_open"} and row is not None:
                cell = row.cells[column]
                column += 1
                paragraph = cell.paragraphs[0]
            elif kind in {"heading_open", "paragraph_open"}:
                if cell is not None:
                    continue
                if kind == "heading_open":
                    paragraph = document.add_heading(level=min(int(token.tag[1:]), 9))
                else:
                    style = None
                    if list_stack:
                        style = (
                            "List Bullet"
                            if list_stack[-1] == "bullet_list_open"
                            else "List Number"
                        )
                    paragraph = document.add_paragraph(style=style)
            elif kind == "inline" and paragraph is not None:
                _append_inline(paragraph, token.children or [])
            elif kind == "fence":
                document.add_paragraph(token.content, style="No Spacing")
            elif kind in {"th_close", "td_close"}:
                cell = None
                paragraph = None
            elif kind in {"heading_close", "paragraph_close"}:
                if cell is None:
                    paragraph = None
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()


def _export_reader_markdown(reader: ReaderReport) -> str:
    """Link source URLs and remove UI-only evidence targets in binary exports."""
    markdown = reader.markdown
    if reader.bibliography_span:
        start, end = reader.bibliography_span

        def autolink(match: re.Match) -> str:
            value = match.group()
            url = value.rstrip(".,;。；，)")
            return f"<{url}>" + value[len(url) :]

        bibliography = re.sub(r"https?://[^\s<>]+", autolink, markdown[start:end])
        markdown = markdown[:start] + bibliography + markdown[end:]
    # These targets represent browser evidence buttons, not document bookmarks.
    for link in reader.citation_links:
        markdown = markdown.replace(
            f"[{link['label']}](#{link['anchor']})", link["label"]
        )
    return markdown


def _append_inline(paragraph, children: list) -> None:
    """Map Markdown inline tokens to native Word runs and safe hyperlinks."""
    bold = italic = False
    link: str | None = None
    for child in children:
        kind = child.type
        if kind == "strong_open":
            bold = True
        elif kind == "strong_close":
            bold = False
        elif kind == "em_open":
            italic = True
        elif kind == "em_close":
            italic = False
        elif kind == "link_open":
            link = child.attrGet("href")
        elif kind == "link_close":
            link = None
        elif kind in {"text", "code_inline"}:
            if link and re.match(r"^(https?://|mailto:)", link):
                _add_hyperlink(paragraph, child.content, link)
            else:
                run = paragraph.add_run(child.content)
                run.bold = bold
                run.italic = italic
                if kind == "code_inline":
                    run.font.name = "Consolas"
        elif kind in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        elif kind == "image":
            paragraph.add_run("[image omitted]")


def _add_hyperlink(paragraph, label: str, url: str) -> None:
    """Add a clickable link without dereferencing its untrusted target."""
    relation = paragraph.part.relate_to(
        url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
